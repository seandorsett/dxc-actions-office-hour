#!/usr/bin/env python3
"""Check slide count, speaker-note fidelity, timing, template size, and bounds.

Usage: python scripts/check_deck.py
Prerequisites: pip install -r requirements-presentation.txt
These checks do not replace visual review or live GitHub demo rehearsal.
"""

import json
from pathlib import Path
import re
from xml.etree import ElementTree
from zipfile import ZipFile

from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]


def check():
    plan = json.loads((ROOT / "presentation" / "slide-content.json").read_text(encoding="utf-8"))
    output = ROOT / "outputs" / "slides" / "dxc-actions-office-hour-2026-10-13.pptx"
    deck = Presentation(output)
    template = Presentation(ROOT / plan["template"])
    errors = []
    if len(deck.slides) != len(plan["slides"]):
        errors.append("Slide count differs from plan.")
    if (deck.slide_width, deck.slide_height) != (template.slide_width, template.slide_height):
        errors.append("Template slide dimensions changed.")
    elapsed = 0
    for index, (slide, item) in enumerate(zip(deck.slides, plan["slides"]), 1):
        if item["kind"] == "qa" and elapsed != 50:
            errors.append(f"Q&A starts at {elapsed}, not minute 50.")
        expected_footer = ("Appendix" if item["minutes"] == 0 else
                           f"{elapsed:02d}-{elapsed + item['minutes']:02d} min")
        if not any(shape.has_text_frame and shape.text == expected_footer
                   for shape in slide.shapes):
            errors.append(f"Slide {index}: timing footer differs from the plan.")
        elapsed += item["minutes"]
        if slide.notes_slide.notes_text_frame.text != item["notes"]:
            errors.append(f"Slide {index}: notes differ from full transcript.")
        if len(item["notes"].split()) < 60:
            errors.append(f"Slide {index}: presenter transcript unexpectedly short.")
        for shape in slide.shapes:
            if (shape.left < 0 or shape.top < 0
                    or shape.left + shape.width > deck.slide_width
                    or shape.top + shape.height > deck.slide_height):
                errors.append(f"Slide {index}: shape outside template dimensions.")
            if shape.has_text_frame and re.search(
                    r"\[Speaker Name\]|\[Team / Group Name\]|lorem ipsum|TODO|FIXME",
                    shape.text, re.I):
                errors.append(f"Slide {index}: template example text remains.")
    if elapsed != 60:
        errors.append(f"Planned session lasts {elapsed}, not 60 minutes.")
    with ZipFile(output) as archive:
        master = archive.read("ppt/slideMasters/slideMaster1.xml").decode("utf-8")
        if "DXC" not in master:
            errors.append("DXC branding missing from template master.")
        for file in archive.namelist():
            if file.startswith("ppt/slides/_rels/") and file.endswith(".rels"):
                relationships = ElementTree.fromstring(archive.read(file))
                for relationship in relationships:
                    if (relationship.get("TargetMode") == "External"
                            and not relationship.get("Target", "").startswith("https://")):
                        errors.append(f"Unexpected external relationship in {file}.")
    if errors:
        raise RuntimeError("\n".join(errors))
    print(f"Deck checks passed: {len(deck.slides)} slides; original DXC template; "
          "verbatim notes; 50-minute handoff; 10-minute Q&A.")


if __name__ == "__main__":
    check()
