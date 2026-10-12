#!/usr/bin/env python3
"""Build the Office Hour deck from the supplied DXC template.

Usage: python scripts/generate_deck.py
Prerequisites: pip install -r requirements-presentation.txt
The actual template masters/layouts/assets are retained. No external skill
installation, network connection, or credentials are needed in GitHub Actions.
"""

from __future__ import annotations

import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs" / "slides" / "dxc-actions-office-hour-2026-10-13.pptx"
INK = RGBColor.from_string("0E1022")
BLUE = RGBColor.from_string("345282")
GRAY = RGBColor.from_string("555B65")
PALE = RGBColor.from_string("F2F4F7")
WHITE = RGBColor.from_string("FFFFFF")
FONT = "Arial"


def box(slide, text, x, y, w, h, size=20, color=INK, bold=False,
        fill=None, align=PP_ALIGN.LEFT, font=FONT):
    if fill is None:
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    else:
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                      Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
        shape.line.fill.background()
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    frame.margin_left = frame.margin_right = Inches(0.16 if fill else 0.02)
    frame.margin_top = frame.margin_bottom = Inches(0.13 if fill else 0.02)
    frame.vertical_anchor = MSO_ANCHOR.TOP
    for index, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.space_after = Pt(0 if font == "Consolas" else 8)
        if font == "Consolas":
            paragraph.line_spacing = 1.1
        run = paragraph.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return shape


def card(slide, title, body, x, y, w, h, dark=False):
    panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                  Inches(x), Inches(y), Inches(w), Inches(h))
    panel.fill.solid()
    panel.fill.fore_color.rgb = INK if dark else PALE
    panel.line.fill.background()
    color = WHITE if dark else INK
    if title:
        box(slide, title, x + 0.22, y + 0.20, w - 0.44, 0.58, 23, color, True)
    if body:
        box(slide, body, x + 0.22, y + 1.02, w - 0.44, h - 1.24,
            20, WHITE if dark else GRAY)


def clean_slide(prs, layout_name):
    layout = next(layout for layout in prs.slide_layouts if layout.name == layout_name)
    slide = prs.slides.add_slide(layout)
    for shape in list(slide.shapes):
        if shape.is_placeholder:
            shape._element.getparent().remove(shape._element)
    return slide


def footer(slide, data, elapsed, index):
    if source := data.get("source"):
        shape = box(slide, source[0], 0.65, 8.10, 10.5, 0.25, 11, BLUE)
        shape.text_frame.paragraphs[0].runs[0].hyperlink.address = source[1]
    label = ("Appendix" if data["minutes"] == 0 else
             f"{elapsed:02d}-{elapsed + data['minutes']:02d} min")
    box(slide, label, 12.00, 8.10, 3.35, 0.25, 11, GRAY, align=PP_ALIGN.RIGHT)


def table(slide, headers, rows, agenda=False):
    all_rows = [headers] + rows
    height = 4.75 if agenda else min(4.10, 0.78 * len(all_rows))
    shape = slide.shapes.add_table(len(all_rows), len(headers), Inches(0.65),
                                   Inches(2.95), Inches(14.7), Inches(height))
    grid = shape.table
    widths = ([1.55, 4.65, 8.5] if agenda else
              [3.25, 6.20, 5.25] if len(headers) == 3 else [2.9, 2.4, 2.4, 7.0])
    for column, width in zip(grid.columns, widths):
        column.width = Inches(width)
    for row_index, cells in enumerate(all_rows):
        for col_index, value in enumerate(cells):
            cell = grid.cell(row_index, col_index)
            cell.fill.solid()
            cell.fill.fore_color.rgb = INK if row_index == 0 else (PALE if row_index % 2 else WHITE)
            cell.margin_left = cell.margin_right = Inches(0.15)
            cell.margin_top = cell.margin_bottom = Inches(0.065)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            frame = cell.text_frame
            frame.clear()
            frame.word_wrap = True
            frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
            frame.paragraphs[0].line_spacing = 1.0
            run = frame.paragraphs[0].add_run()
            run.text = value
            run.font.name = FONT
            run.font.size = Pt(18)
            run.font.bold = row_index == 0
            run.font.color.rgb = WHITE if row_index == 0 else INK


def content(prs, data, index, elapsed):
    kind = data["kind"]
    if kind == "cover":
        slide = clean_slide(prs, "Title 4")
        box(slide, data["body"], 6.53, 5.88, 8.4, 0.36, 12, GRAY)
        box(slide, data["title"], 6.53, 6.25, 8.5, 1.18, 34, INK, True)
        box(slide, data["subtitle"], 6.53, 7.60, 8.5, 0.40, 15, INK)
    elif kind == "qa":
        slide = clean_slide(prs, "Section Slide")
        box(slide, data["title"], 0.53, 3.48, 11.7, 1.0, 54, INK, True)
        box(slide, data["subtitle"], 0.60, 5.14, 11.3, 0.65, 24)
        box(slide, data["body"], 0.60, 6.12, 11.3, 0.92, 20, GRAY)
    else:
        slide = clean_slide(prs, "Title and Content")
        box(slide, data.get("section", "DXC | GitHub Actions Office Hour"),
            0.65, 0.97, 14.7, 0.26, 11, GRAY)
        box(slide, data["title"], 0.53, 1.38, 14.95, 1.08, 32, INK, True)
        if kind == "agenda":
            table(slide, ["Minutes", "Segment", "Outcome"], data["rows"], agenda=True)
        elif kind == "table":
            table(slide, data["headers"], data["rows"])
        elif kind in ("cards", "columns"):
            count = len(data["items"])
            gap = 0.30
            width = (14.7 - gap * (count - 1)) / count
            for i, (title, body) in enumerate(data["items"]):
                card(slide, title, body, 0.65 + i * (width + gap), 3.10, width, 3.65,
                     dark=kind == "columns" and i == 1)
        elif kind == "code":
            box(slide, data["code"], 0.65, 3.02, 9.2, 3.95, 15, INK,
                fill=PALE, font="Consolas")
            for i, (title, body) in enumerate(data["items"]):
                box(slide, title, 10.2, 3.10 + i * 1.30, 5.1, 0.35, 20, INK, True)
                box(slide, body, 10.2, 3.57 + i * 1.30, 5.1, 0.67, 19, GRAY)
        elif kind == "demo":
            box(slide, data["workflow"], 0.65, 2.62, 14.7, 0.26, 13, BLUE, font="Consolas")
            for i, (title, body) in enumerate(data["items"]):
                y = 3.06 + i * 1.43
                panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                              Inches(0.65), Inches(y), Inches(10.75), Inches(1.18))
                panel.fill.solid()
                panel.fill.fore_color.rgb = PALE
                panel.line.fill.background()
                box(slide, f"{i + 1}. {title}", 0.90, y + 0.14, 3.0, 0.82, 21, INK, True)
                box(slide, body, 4.00, y + 0.14, 7.15, 0.89, 20, GRAY)
            card(slide, "Run from GitHub", "Actions tab\nManual trigger\nJob summaries\nEvidence artifacts",
                 11.72, 3.06, 3.63, 4.05, dark=True)
        elif kind == "flow":
            count = len(data["items"])
            width = (14.7 - 0.45 * (count - 1)) / count
            positions = [0.65 + i * (width + 0.45) for i in range(count)]
            for x in positions:
                panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                              Inches(x), Inches(3.70), Inches(width), Inches(2.7))
                panel.fill.solid()
                panel.fill.fore_color.rgb = PALE
                panel.line.fill.background()
            for x in positions[:-1]:
                arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,
                                              Inches(x + width + 0.075), Inches(4.72),
                                              Inches(0.30), Inches(0.38))
                arrow.fill.solid()
                arrow.fill.fore_color.rgb = BLUE
                arrow.line.fill.background()
            for x, (title, body) in zip(positions, data["items"]):
                box(slide, title, x + 0.20, 3.95, width - 0.40, 0.55, 23, INK, True)
                box(slide, body, x + 0.20, 4.92, width - 0.40, 1.13, 19, GRAY)
        elif kind == "metrics":
            for i, (value, title, body) in enumerate(data["items"]):
                x = 0.65 + i * 7.50
                card(slide, "", "", x, 3.08, 7.20, 3.72)
                box(slide, value, x + 0.28, 3.35, 6.60, 1.02, 54, BLUE, True)
                box(slide, title, x + 0.28, 4.64, 6.60, 0.52, 24, INK, True)
                box(slide, body, x + 0.28, 5.55, 6.60, 0.92, 20, GRAY)
        elif kind == "links":
            for i, (title, url) in enumerate(data["links"]):
                y = 2.93 + i * 0.73
                shape = box(slide, title, 0.85, y, 14.1, 0.48, 22, BLUE, True)
                shape.text_frame.paragraphs[0].runs[0].hyperlink.address = url
        else:
            raise ValueError(f"Unsupported layout: {kind}")
        if callout := data.get("callout"):
            box(slide, callout, 0.65, 7.38, 14.70, 0.60, 17, INK, fill=PALE)
    slide.notes_slide.notes_text_frame.text = data["notes"]
    footer(slide, data, elapsed, index)
    return slide


def build():
    plan = json.loads((ROOT / "presentation" / "slide-content.json").read_text(encoding="utf-8"))
    prs = Presentation(ROOT / plan["template"])
    for slide_id in list(prs.slides._sldIdLst):
        prs.part.drop_rel(slide_id.rId)
        prs.slides._sldIdLst.remove(slide_id)
    elapsed = 0
    for index, data in enumerate(plan["slides"], 1):
        content(prs, data, index, elapsed)
        elapsed += data["minutes"]
    prs.core_properties.title = plan["title"]
    prs.core_properties.subject = "Five GitHub Actions updates | 50-minute session plus 10-minute Q&A"
    prs.core_properties.author = "DXC GitHub Actions Office Hour"
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"Created {OUT.name}: {len(prs.slides)} slides with full speaker notes.")


if __name__ == "__main__":
    build()
